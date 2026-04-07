"""
Servicio de Importación Masiva de Empleados — Excel (.xlsx)
"""

import io
import logging
import secrets
import string
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models.team import Team
from app.models.user import User, UserRole
from app.models.vacation_balance import VacationBalance
from app.repositories.audit_repo import AuditRepository
from app.repositories.team_repo import TeamRepository
from app.repositories.user_repo import UserRepository
from app.repositories.vacation_balance_repo import VacationBalanceRepository

logger = logging.getLogger(__name__)

# ── Colores ────────────────────────────────────────────────────────────────
_AZUL_OSCURO = "1B2A4A"
_AZUL_MEDIO = "2E5090"
_AZUL_CLARO = "D6E4F0"
_VERDE = "548235"
_VERDE_CLARO = "E2EFDA"
_GRIS = "F2F2F2"
_BORDE = "B4B4B4"
_BLANCO = "FFFFFF"
_ROJO = "C00000"
_ROJO_CLARO = "FDE8E8"
_NARANJA = "ED7D31"
_AMARILLO = "FFF2CC"

# ── LFT México 2023 ───────────────────────────────────────────────────────
LFT_TABLE = [
    (1,1,12),(2,2,14),(3,3,16),(4,4,18),(5,5,20),
    (6,10,22),(11,15,24),(16,20,26),(21,25,28),(26,30,30),(31,35,32),
]

MAX_ROWS = 500
REQUIRED_COLUMNS = ["nombre_completo","email","rol","equipo","fecha_ingreso"]
OPTIONAL_COLUMNS = ["puesto","telefono","contacto_emergencia","gerente_email"]
ALL_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
VALID_ROLES = {"EMPLOYEE","MANAGER"}

ROLE_MAP = {
    "empleado":"EMPLOYEE","employee":"EMPLOYEE",
    "gerente":"MANAGER","manager":"MANAGER","jefe":"MANAGER",
    "líder":"MANAGER","lider":"MANAGER","supervisor":"MANAGER",
}


def compute_vacation_days(hire_date: date, reference_date: date | None = None) -> int:
    ref = reference_date or date.today()
    if hire_date > ref:
        return 0
    years = ref.year - hire_date.year
    if (ref.month, ref.day) < (hire_date.month, hire_date.day):
        years -= 1
    years = max(years, 0)
    if years < 1:
        days_worked = (ref - hire_date).days
        return max(int(round(12 * days_worked / 365, 2)), 0)
    for s, e, d in LFT_TABLE:
        if s <= years <= e:
            return d
    return 32


def _gen_pwd(length: int = 12) -> str:
    alpha = string.ascii_letters + string.digits + "!@#$%"
    while True:
        p = "".join(secrets.choice(alpha) for _ in range(length))
        if any(c.isupper() for c in p) and any(c.islower() for c in p) and any(c.isdigit() for c in p) and any(c in "!@#$%" for c in p):
            return p


def _normalize_role(raw: str) -> str | None:
    k = raw.strip().lower()
    if k in ROLE_MAP:
        return ROLE_MAP[k]
    u = raw.strip().upper()
    return u if u in VALID_ROLES else None


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d","%d/%m/%Y","%d-%m-%Y","%m/%d/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _resolve_team(name: str, all_teams: list[Team]) -> tuple[Team | None, str | None]:
    clean = name.strip()
    key = clean.lower()
    for t in all_teams:
        if t.name == clean:
            return t, None
    for t in all_teams:
        if t.name.lower() == key:
            return t, None
    for t in all_teams:
        if key in t.name.lower() or t.name.lower() in key:
            return t, None
    avail = ", ".join(t.name for t in all_teams) or "No hay equipos"
    return None, avail


def _brd(color: str = _BORDE) -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


# ═══════════════════════════════════════════════════════════════════════════
# PLANTILLA
# ═══════════════════════════════════════════════════════════════════════════

def generate_template(db: Session) -> bytes:
    team_repo = TeamRepository(db)
    user_repo = UserRepository(db)
    teams = team_repo.list_all()
    team_names = [t.name for t in teams]
    managers = user_repo.list_all(role="MANAGER")
    manager_info = [(m.full_name, m.email) for m in managers]

    wb = Workbook()
    bd = _brd()
    hdr_fill = PatternFill(start_color=_AZUL_OSCURO, end_color=_AZUL_OSCURO, fill_type="solid")
    hdr_font = Font(name="Calibri", size=11, bold=True, color=_BLANCO)
    sec_fill = PatternFill(start_color=_AZUL_MEDIO, end_color=_AZUL_MEDIO, fill_type="solid")
    sec_font = Font(name="Calibri", size=13, bold=True, color=_BLANCO)
    th_fill = PatternFill(start_color=_AZUL_CLARO, end_color=_AZUL_CLARO, fill_type="solid")
    th_font = Font(name="Calibri", size=10, bold=True, color=_AZUL_OSCURO)
    tf = Font(name="Calibri", size=10, color="333333")
    alt = PatternFill(start_color=_GRIS, end_color=_GRIS, fill_type="solid")

    # ── HOJA 1: GUÍA PARA RH ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Guia para RH"
    ws.sheet_properties.tabColor = _AZUL_OSCURO
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 20

    for c in range(1, 6):
        ws.cell(row=1, column=c).fill = hdr_fill
        ws.cell(row=2, column=c).fill = hdr_fill
    ws.merge_cells("B1:E1")
    ws.merge_cells("B2:E2")
    ws.cell(row=1, column=2, value="SEEKOP — Importación Masiva de Empleados").font = Font(name="Calibri", size=18, bold=True, color=_BLANCO)
    ws.cell(row=1, column=2).alignment = Alignment(vertical="center")
    ws.cell(row=2, column=2, value="Guía completa para Recursos Humanos").font = Font(name="Calibri", size=12, color="A0B4D0", italic=True)
    ws.row_dimensions[1].height = 35
    ws.row_dimensions[2].height = 25

    r = 4

    def seccion(txt):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        for cc in range(2, 6):
            ws.cell(row=r, column=cc).fill = sec_fill
        ws.cell(row=r, column=2, value=txt).font = sec_font
        ws.row_dimensions[r].height = 30
        r += 1

    def paso(num, txt):
        nonlocal r
        ws.cell(row=r, column=2, value=num).font = Font(name="Calibri", size=11, bold=True, color=_AZUL_MEDIO)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        c = ws.cell(row=r, column=3, value=txt)
        c.font = tf
        c.alignment = Alignment(wrap_text=True)
        r += 1

    seccion("PASOS PARA LA IMPORTACIÓN")
    paso("Paso 1:", "Abre la hoja 'Empleados' de este archivo.")
    paso("Paso 2:", "Llena los datos de cada empleado. Elimina o sobrescribe la fila de ejemplo.")
    paso("Paso 3:", "Guarda el archivo y súbelo desde el panel de Importación Masiva en el sistema.")
    paso("Paso 4:", "El sistema validará, creará cuentas y calculará vacaciones automáticamente.")
    paso("Paso 5:", "Descarga el archivo de resultados con las contraseñas temporales.")
    r += 1

    seccion("CAMPOS DEL ARCHIVO")
    r += 1
    for ci, h in enumerate(["Campo", "Descripción", "Ejemplo", "Obligatorio"], start=2):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = th_font
        c.fill = th_fill
        c.border = bd
        c.alignment = Alignment(horizontal="center")
    r += 1

    campos = [
        ("nombre_completo", "Nombre y apellidos del empleado", "María López García", "SÍ"),
        ("email", "Correo electrónico corporativo (único)", "maria.lopez@seekop.com", "SÍ"),
        ("rol", "Empleado o Gerente (ver sección Roles)", "Empleado", "SÍ"),
        ("equipo", "Nombre del equipo (ver hoja Equipos)", team_names[0] if team_names else "—", "SÍ"),
        ("fecha_ingreso", "Fecha de inicio laboral", "2023-03-15", "SÍ"),
        ("puesto", "Título o cargo del empleado", "Analista de Datos", "No"),
        ("telefono", "Número de teléfono", "5512345678", "No"),
        ("contacto_emergencia", "Nombre y teléfono de contacto", "Juan López — 5598765432", "No"),
        ("gerente_email", "Email del gerente asignado (ver hoja Equipos)", manager_info[0][1] if manager_info else "gerente@seekop.com", "No"),
    ]
    for i, (ca, de, ej, ob) in enumerate(campos):
        fl = alt if i % 2 == 0 else PatternFill(fill_type=None)
        for ci, v in enumerate([ca, de, ej, ob], start=2):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = Font(name="Calibri", size=10, bold=(ci == 5 and v == "SÍ"), color=_ROJO if (ci == 5 and v == "SÍ") else (_VERDE if (ci == 5 and v == "No") else "333333"))
            c.border = bd
            c.fill = fl
            c.alignment = Alignment(wrap_text=True, vertical="center")
        r += 1
    r += 1

    seccion("ROLES VÁLIDOS")
    r += 1
    for ci, h in enumerate(["Escribir en el Excel", "Alternativas aceptadas", "Descripción"], start=2):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = th_font
        c.fill = th_fill
        c.border = bd
        c.alignment = Alignment(horizontal="center")
    r += 1
    for i, (ro, al, de) in enumerate([
        ("Empleado", "empleado, Employee", "Personal operativo sin subordinados directos"),
        ("Gerente", "gerente, Manager, jefe, supervisor, líder", "Personal con equipo a su cargo"),
    ]):
        fl = alt if i % 2 == 0 else PatternFill(fill_type=None)
        for ci, v in enumerate([ro, al, de], start=2):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = tf
            c.border = bd
            c.fill = fl
        r += 1
    r += 1

    seccion("FORMATOS DE FECHA ACEPTADOS")
    r += 1
    for ci, h in enumerate(["Formato", "Ejemplo"], start=2):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = th_font
        c.fill = th_fill
        c.border = bd
        c.alignment = Alignment(horizontal="center")
    r += 1
    for fn, fe in [("AAAA-MM-DD (recomendado)", "2023-03-15"), ("DD/MM/AAAA", "15/03/2023"), ("DD-MM-AAAA", "15-03-2023")]:
        ws.cell(row=r, column=2, value=fn).font = tf
        ws.cell(row=r, column=2).border = bd
        ws.cell(row=r, column=3, value=fe).font = tf
        ws.cell(row=r, column=3).border = bd
        r += 1
    r += 1

    seccion("CÁLCULO DE VACACIONES — LFT MÉXICO 2023")
    r += 1
    for ci, h in enumerate(["Antigüedad", "Días de vacaciones"], start=2):
        c = ws.cell(row=r, column=ci, value=h)
        c.font = th_font
        c.fill = th_fill
        c.border = bd
        c.alignment = Alignment(horizontal="center")
    r += 1
    for i, (an, di) in enumerate([
        ("1 año cumplido", 12), ("2 años", 14), ("3 años", 16), ("4 años", 18), ("5 años", 20),
        ("6 a 10 años", 22), ("11 a 15 años", 24), ("16 a 20 años", 26),
        ("21 a 25 años", 28), ("26 a 30 años", 30), ("31 a 35 años", 32),
    ]):
        fl = alt if i % 2 == 0 else PatternFill(fill_type=None)
        c1 = ws.cell(row=r, column=2, value=an)
        c1.font = tf
        c1.border = bd
        c1.fill = fl
        c2 = ws.cell(row=r, column=3, value=f"{di} días")
        c2.font = Font(name="Calibri", size=10, bold=True, color=_VERDE)
        c2.border = bd
        c2.fill = fl
        c2.alignment = Alignment(horizontal="center")
        r += 1
    r += 1

    nota_fill = PatternFill(start_color=_AMARILLO, end_color=_AMARILLO, fill_type="solid")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for cc in range(2, 6):
        ws.cell(row=r, column=cc).fill = nota_fill
    ws.cell(row=r, column=2, value="Nota: Si el empleado aún no cumple 1 año, se le asigna la parte proporcional de 12 días.").font = Font(name="Calibri", size=10, italic=True, color=_NARANJA)
    r += 2

    seccion("CONTRASEÑAS TEMPORALES")
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r+2, end_column=5)
    ws.cell(row=r, column=2, value=(
        "Al importar, el sistema genera una contraseña temporal única por empleado. "
        "Al finalizar, descarga el Excel de resultados con las contraseñas. "
        "Cada empleado deberá cambiar su contraseña al primer inicio de sesión.\n\n"
        "IMPORTANTE: El archivo de contraseñas es confidencial. Compártelo solo con personal autorizado."
    )).font = tf
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    r += 4

    seccion("REGLAS IMPORTANTES")
    r += 1
    for regla in [
        "No modifiques los encabezados de la hoja 'Empleados'.",
        "Máximo 500 empleados por archivo.",
        "Los correos deben ser únicos (no repetidos en el archivo ni en el sistema).",
        "El equipo debe coincidir con uno registrado (ver hoja 'Equipos Disponibles').",
        "Si hay errores en una fila, esa fila se salta y las demás se procesan normalmente.",
    ]:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.cell(row=r, column=2, value=f"  •  {regla}").font = tf
        ws.row_dimensions[r].height = 20
        r += 1

    # ── HOJA 2: EMPLEADOS (rediseño profesional) ─────────────────────────
    ws2 = wb.create_sheet("Empleados")
    ws2.sheet_properties.tabColor = _VERDE

    # Columnas: A=número fila, B-J=datos
    COL_NUM = 1       # A — # fila
    COL_START = 2     # B — primer campo de datos
    COL_END = 10      # J — último campo de datos
    TOTAL_COLS = COL_END

    col_widths = {1: 5, 2: 34, 3: 38, 4: 18, 5: 30, 6: 20, 7: 28, 8: 20, 9: 34, 10: 36}
    for ci, w in col_widths.items():
        ws2.column_dimensions[get_column_letter(ci)].width = w

    # ─── Fila 1: Banner principal ───
    banner_fill = PatternFill(start_color=_AZUL_OSCURO, end_color=_AZUL_OSCURO, fill_type="solid")
    for ci in range(1, TOTAL_COLS + 1):
        ws2.cell(row=1, column=ci).fill = banner_fill
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COLS)
    b1 = ws2.cell(row=1, column=1, value="SEEKOP — Registro Masivo de Empleados")
    b1.font = Font(name="Calibri", size=16, bold=True, color=_BLANCO)
    b1.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 40

    # ─── Fila 2: Subtítulo / instrucción rápida ───
    sub_fill = PatternFill(start_color=_AZUL_MEDIO, end_color=_AZUL_MEDIO, fill_type="solid")
    for ci in range(1, TOTAL_COLS + 1):
        ws2.cell(row=2, column=ci).fill = sub_fill
    ws2.merge_cells(start_row=2, start_column=1, end_row=2, end_column=TOTAL_COLS)
    b2 = ws2.cell(row=2, column=1,
        value="Completa los datos de cada empleado en las columnas de abajo. Los campos marcados con ★ son obligatorios. "
              "Las columnas Rol y Equipo tienen lista desplegable.")
    b2.font = Font(name="Calibri", size=10, color="D0DCF0", italic=True)
    b2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws2.row_dimensions[2].height = 30

    # ─── Fila 3: Separador con leyenda ───
    legend_fill = PatternFill(start_color=_AMARILLO, end_color=_AMARILLO, fill_type="solid")
    for ci in range(1, TOTAL_COLS + 1):
        ws2.cell(row=3, column=ci).fill = legend_fill
    ws2.merge_cells(start_row=3, start_column=1, end_row=3, end_column=TOTAL_COLS)
    lg = ws2.cell(row=3, column=1,
        value="★ = Obligatorio     ○ = Opcional     |     Roles: Empleado, Gerente     |     "
              f"Equipos: {', '.join(team_names) if team_names else 'Ver hoja Equipos Disponibles'}")
    lg.font = Font(name="Calibri", size=9, bold=True, color="6B5B00")
    lg.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[3].height = 22

    # ─── Fila 4: Etiquetas obligatorio / opcional ───
    headers_meta = [
        ("",      ""),
        ("nombre_completo",    "★ OBLIGATORIO"),
        ("email",              "★ OBLIGATORIO"),
        ("rol",                "★ OBLIGATORIO"),
        ("equipo",             "★ OBLIGATORIO"),
        ("fecha_ingreso",      "★ OBLIGATORIO"),
        ("puesto",             "○ Opcional"),
        ("telefono",           "○ Opcional"),
        ("contacto_emergencia","○ Opcional"),
        ("gerente_email",      "○ Opcional"),
    ]
    tag_fill_req = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
    tag_fill_opt = PatternFill(start_color=_VERDE_CLARO, end_color=_VERDE_CLARO, fill_type="solid")
    for ci, (_, tag) in enumerate(headers_meta, start=1):
        if ci == 1:
            ws2.cell(row=4, column=ci).fill = PatternFill(start_color=_GRIS, end_color=_GRIS, fill_type="solid")
            continue
        is_req = tag.startswith("★")
        c = ws2.cell(row=4, column=ci, value=tag)
        c.font = Font(name="Calibri", size=8, bold=True, color=_ROJO if is_req else _VERDE)
        c.fill = tag_fill_req if is_req else tag_fill_opt
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = bd
    ws2.row_dimensions[4].height = 16

    # ─── Fila 5: Encabezados principales (nombres legibles en español) ───
    header_labels = [
        "#",
        "Nombre Completo",
        "Correo Electrónico",
        "Rol",
        "Equipo",
        "Fecha de Ingreso",
        "Puesto / Cargo",
        "Teléfono",
        "Contacto de Emergencia",
        "Email del Gerente",
    ]
    for ci, label in enumerate(header_labels, start=1):
        c = ws2.cell(row=5, column=ci, value=label)
        c.font = Font(name="Calibri", size=11, bold=True, color=_BLANCO)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _brd("5B6E8F")
    ws2.row_dimensions[5].height = 32

    # ─── Fila 6: Descripción / ayuda debajo de cada encabezado ───
    help_texts = [
        "",
        "Nombre y apellidos",
        "nombre@empresa.com",
        "Seleccionar ▼",
        "Seleccionar ▼",
        "AAAA-MM-DD",
        "Título del cargo",
        "10 dígitos",
        "Nombre — Teléfono",
        "Ver hoja Equipos ▼",
    ]
    help_fill = PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid")
    for ci, txt in enumerate(help_texts, start=1):
        c = ws2.cell(row=6, column=ci, value=txt)
        c.font = Font(name="Calibri", size=9, color="7F8FA6", italic=True)
        c.fill = help_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _brd("D5DDE5")
    ws2.row_dimensions[6].height = 20

    # ─── Fila 7: Encabezados técnicos (nombres de columna reales, discretos) ───
    tech_headers = ["", "nombre_completo", "email", "rol", "equipo", "fecha_ingreso", "puesto", "telefono", "contacto_emergencia", "gerente_email"]
    tech_fill = PatternFill(start_color="D9E2EC", end_color="D9E2EC", fill_type="solid")
    for ci, th in enumerate(tech_headers, start=1):
        c = ws2.cell(row=7, column=ci, value=th)
        c.font = Font(name="Consolas", size=8, color="8899AA")
        c.fill = tech_fill
        c.alignment = Alignment(horizontal="center")
        c.border = _brd("C5D0DB")
    ws2.row_dimensions[7].height = 16

    # ─── Fila 8: Ejemplo pre-llenado ───
    DATA_START_ROW = 8
    example_fill_row = PatternFill(start_color=_VERDE_CLARO, end_color=_VERDE_CLARO, fill_type="solid")
    example_vals = [
        "1",
        "María López García",
        "maria.lopez@seekop.com",
        "Empleado",
        team_names[0] if team_names else "—",
        "2023-03-15",
        "Analista de Datos",
        "5512345678",
        "Juan López — 5598765432",
        manager_info[0][1] if manager_info else "gerente@seekop.com",
    ]
    for ci, v in enumerate(example_vals, start=1):
        c = ws2.cell(row=DATA_START_ROW, column=ci, value=v)
        c.font = Font(name="Calibri", size=10, color="70876E", italic=True)
        c.fill = example_fill_row
        c.border = bd
        c.alignment = Alignment(horizontal="left" if ci > 1 else "center", vertical="center")
    ws2.row_dimensions[DATA_START_ROW].height = 24

    # ─── Filas 9+: Área de datos vacía con formato alterno pre-aplicado ───
    row_fill_a = PatternFill(start_color=_BLANCO, end_color=_BLANCO, fill_type="solid")
    row_fill_b = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")
    light_brd = _brd("E0E5EB")
    for ri in range(DATA_START_ROW + 1, DATA_START_ROW + 31):  # 30 filas pre-formateadas
        fill = row_fill_a if (ri - DATA_START_ROW) % 2 != 0 else row_fill_b
        ws2.cell(row=ri, column=1, value=ri - DATA_START_ROW + 1).font = Font(name="Calibri", size=9, color="AABBCC")
        ws2.cell(row=ri, column=1).alignment = Alignment(horizontal="center")
        ws2.cell(row=ri, column=1).fill = PatternFill(start_color="EEF1F5", end_color="EEF1F5", fill_type="solid")
        for ci in range(COL_START, TOTAL_COLS + 1):
            c = ws2.cell(row=ri, column=ci)
            c.fill = fill
            c.border = light_brd
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(vertical="center")
        ws2.row_dimensions[ri].height = 22

    # ─── Nota al final del área pre-formateada ───
    note_row = DATA_START_ROW + 32
    ws2.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=TOTAL_COLS)
    nc = ws2.cell(row=note_row, column=1,
        value="Puedes agregar más filas debajo. Máximo 500 empleados por archivo. "
              "No modifiques las filas 1-7 (encabezados).")
    nc.font = Font(name="Calibri", size=9, color=_NARANJA, italic=True)
    nc.alignment = Alignment(horizontal="center")

    # ─── Validaciones de datos (dropdowns) ───
    dv_rol = DataValidation(type="list", formula1='"Empleado,Gerente"', allow_blank=False)
    dv_rol.error = "Rol no válido. Seleccione: Empleado o Gerente"
    dv_rol.errorTitle = "Rol inválido"
    dv_rol.prompt = "Seleccione el rol del empleado"
    dv_rol.promptTitle = "Rol"
    dv_rol.showInputMessage = True
    dv_rol.showErrorMessage = True
    ws2.add_data_validation(dv_rol)
    dv_rol.add(f"D{DATA_START_ROW}:D{MAX_ROWS + DATA_START_ROW}")

    if team_names:
        dv_eq = DataValidation(type="list", formula1=f'"{",".join(team_names)}"', allow_blank=False)
        dv_eq.error = "Equipo no encontrado. Seleccione uno de la lista."
        dv_eq.errorTitle = "Equipo inválido"
        dv_eq.prompt = "Seleccione el equipo asignado"
        dv_eq.promptTitle = "Equipo"
        dv_eq.showInputMessage = True
        dv_eq.showErrorMessage = True
        ws2.add_data_validation(dv_eq)
        dv_eq.add(f"E{DATA_START_ROW}:E{MAX_ROWS + DATA_START_ROW}")

    ws2.freeze_panes = f"A{DATA_START_ROW}"
    ws2.sheet_view.showGridLines = False

    # ── HOJA 3: EQUIPOS Y GERENTES ──────────────────────────────────────────
    ws3 = wb.create_sheet("Equipos Disponibles")
    ws3.sheet_properties.tabColor = _NARANJA
    ws3.column_dimensions["A"].width = 3
    ws3.column_dimensions["B"].width = 8
    ws3.column_dimensions["C"].width = 40
    ws3.column_dimensions["D"].width = 5
    ws3.column_dimensions["E"].width = 8
    ws3.column_dimensions["F"].width = 32
    ws3.column_dimensions["G"].width = 38
    ws3.sheet_view.showGridLines = False

    # ── Banner ──
    for cc in range(1, 8):
        ws3.cell(row=1, column=cc).fill = hdr_fill
    ws3.merge_cells("B1:G1")
    ws3.cell(row=1, column=2, value="Equipos y Gerentes registrados en el sistema").font = Font(name="Calibri", size=15, bold=True, color=_BLANCO)
    ws3.cell(row=1, column=2).alignment = Alignment(vertical="center")
    ws3.row_dimensions[1].height = 38
    for cc in range(1, 8):
        ws3.cell(row=2, column=cc).fill = PatternFill(start_color=_AZUL_MEDIO, end_color=_AZUL_MEDIO, fill_type="solid")
    ws3.merge_cells("B2:G2")
    ws3.cell(row=2, column=2, value="Usa estos nombres en las columnas 'equipo' y 'gerente_email' de la hoja Empleados.").font = Font(name="Calibri", size=10, color="D0DCF0", italic=True)
    ws3.cell(row=2, column=2).alignment = Alignment(vertical="center")
    ws3.row_dimensions[2].height = 25

    # ── Sección EQUIPOS (izquierda) ──
    r3 = 4
    ws3.merge_cells("B4:C4")
    for cc in [2, 3]:
        ws3.cell(row=r3, column=cc).fill = sec_fill
    ws3.cell(row=r3, column=2, value="EQUIPOS").font = Font(name="Calibri", size=12, bold=True, color=_BLANCO)
    ws3.cell(row=r3, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[r3].height = 28
    r3 += 1

    for ci, h in enumerate(["#", "Nombre del Equipo"], start=2):
        c = ws3.cell(row=r3, column=ci, value=h)
        c.font = th_font
        c.fill = th_fill
        c.border = bd
        c.alignment = Alignment(horizontal="center")
    r3 += 1

    if team_names:
        for i, tn in enumerate(team_names, start=1):
            fl = alt if i % 2 == 0 else PatternFill(fill_type=None)
            c1 = ws3.cell(row=r3, column=2, value=i)
            c1.font = Font(name="Calibri", size=10, color="999999")
            c1.border = bd
            c1.fill = fl
            c1.alignment = Alignment(horizontal="center")
            c2 = ws3.cell(row=r3, column=3, value=tn)
            c2.font = Font(name="Calibri", size=11, bold=True, color="333333")
            c2.border = bd
            c2.fill = fl
            r3 += 1
    else:
        ws3.cell(row=r3, column=2, value="No hay equipos registrados.").font = Font(name="Calibri", size=10, color=_ROJO, italic=True)
        r3 += 1

    # ── Sección GERENTES (derecha, a la misma altura) ──
    r3m = 4
    ws3.merge_cells("E4:G4")
    for cc in [5, 6, 7]:
        ws3.cell(row=r3m, column=cc).fill = sec_fill
    ws3.cell(row=r3m, column=5, value="GERENTES ACTUALES").font = Font(name="Calibri", size=12, bold=True, color=_BLANCO)
    ws3.cell(row=r3m, column=5).alignment = Alignment(horizontal="center", vertical="center")
    r3m += 1

    for ci, h in enumerate(["#", "Nombre", "Email (usar en gerente_email)"], start=5):
        c = ws3.cell(row=r3m, column=ci, value=h)
        c.font = th_font
        c.fill = th_fill
        c.border = bd
        c.alignment = Alignment(horizontal="center")
    r3m += 1

    if manager_info:
        for i, (mname, memail) in enumerate(manager_info, start=1):
            fl = alt if i % 2 == 0 else PatternFill(fill_type=None)
            cn = ws3.cell(row=r3m, column=5, value=i)
            cn.font = Font(name="Calibri", size=10, color="999999")
            cn.border = bd
            cn.fill = fl
            cn.alignment = Alignment(horizontal="center")
            c_name = ws3.cell(row=r3m, column=6, value=mname)
            c_name.font = Font(name="Calibri", size=11, bold=True, color="333333")
            c_name.border = bd
            c_name.fill = fl
            c_email = ws3.cell(row=r3m, column=7, value=memail)
            c_email.font = Font(name="Consolas", size=10, color=_AZUL_MEDIO)
            c_email.border = bd
            c_email.fill = fl
            r3m += 1
    else:
        ws3.cell(row=r3m, column=5, value="No hay gerentes registrados.").font = Font(name="Calibri", size=10, color=_ROJO, italic=True)
        r3m += 1

    # Nota al final
    note_r = max(r3, r3m) + 1
    ws3.merge_cells(start_row=note_r, start_column=2, end_row=note_r, end_column=7)
    for cc in range(2, 8):
        ws3.cell(row=note_r, column=cc).fill = PatternFill(start_color=_AMARILLO, end_color=_AMARILLO, fill_type="solid")
    ws3.cell(row=note_r, column=2, value="Tip: Copia y pega el email del gerente directamente de esta tabla a la columna 'Email del Gerente' en la hoja Empleados.").font = Font(name="Calibri", size=10, bold=True, color="6B5B00")
    ws3.cell(row=note_r, column=2).alignment = Alignment(horizontal="center", wrap_text=True)

    # ── Dropdown de gerente_email en hoja Empleados ──
    if manager_info:
        mgr_emails = ",".join(e for _, e in manager_info)
        dv_mgr = DataValidation(type="list", formula1=f'"{mgr_emails}"', allow_blank=True)
        dv_mgr.error = "Email de gerente no encontrado. Seleccione de la lista o consulte la hoja Equipos Disponibles."
        dv_mgr.errorTitle = "Gerente no válido"
        dv_mgr.prompt = "Seleccione el email del gerente (opcional)"
        dv_mgr.promptTitle = "Gerente"
        dv_mgr.showInputMessage = True
        dv_mgr.showErrorMessage = True
        ws2.add_data_validation(dv_mgr)
        dv_mgr.add(f"J{DATA_START_ROW}:J{MAX_ROWS + DATA_START_ROW}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# PROCESAR IMPORTACIÓN
# ═══════════════════════════════════════════════════════════════════════════

def process_import(
    file_bytes: bytes, db: Session, actor_user_id: str, actor_role: str,
    *, preview_only: bool = False,
) -> tuple[list[dict], int, int, str]:
    """Procesa importación. Retorna (results, created, errors, batch_id).
    Si preview_only=True, solo valida sin crear usuarios (batch_id será vacío)."""
    import uuid as _uuid
    batch_id = "" if preview_only else str(_uuid.uuid4())

    user_repo = UserRepository(db)
    team_repo = TeamRepository(db)
    balance_repo = VacationBalanceRepository(db)
    audit_repo = AuditRepository(db)
    all_teams = team_repo.list_all()

    def _err(detail: str):
        return [{"row":0,"name":"","email":"","role":"","team":"","status":"ERROR","detail":detail}], 0, 1, ""

    try:
        wb = load_workbook(filename=io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        logger.error("Error al abrir Excel: %s", exc)
        return _err(f"No se pudo abrir el archivo: {exc}")

    ws = wb["Empleados"] if "Empleados" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        return _err("El archivo est\u00e1 vac\u00edo.")

    # Buscar la fila de encabezados técnicos (puede estar en filas 1-15)
    raw_headers = None
    data_start = None
    for i, row in enumerate(rows[:15]):
        cells = [str(h).strip().lower() if h else "" for h in row]
        if "nombre_completo" in cells and "email" in cells:
            raw_headers = cells
            data_start = i + 1  # datos empiezan en la fila siguiente
            break
    if raw_headers is None:
        return _err("No se encontraron los encabezados requeridos (nombre_completo, email, etc.).")

    missing = [c for c in REQUIRED_COLUMNS if c not in raw_headers]
    if missing:
        return _err(f"Faltan columnas obligatorias: {', '.join(missing)}")

    col_idx = {c: raw_headers.index(c) for c in ALL_COLUMNS if c in raw_headers}
    seen: set[str] = set()
    data_rows = rows[data_start:]

    # Saltar fila de ejemplo
    if data_rows:
        e_idx = col_idx.get("email", 1)
        first_email = str(data_rows[0][e_idx] if e_idx < len(data_rows[0]) else "").strip().lower()
        if first_email in ("maria.lopez@seekop.com", "juan.perez@empresa.com"):
            data_rows = data_rows[1:]

    # Saltar fila de nota de ejemplo
    if data_rows:
        first_val = str(data_rows[0][0] if data_rows[0] and data_rows[0][0] else "").strip()
        if first_val.startswith("↑"):
            data_rows = data_rows[1:]

    # Filtrar filas vacías: solo considerar filas que tengan datos en columnas reales
    # (no la columna de # fila ni filas de nota)
    name_idx = col_idx.get("nombre_completo")
    email_idx = col_idx.get("email")
    def _row_has_data(row):
        if not row:
            return False
        for idx in (name_idx, email_idx):
            if idx is not None and idx < len(row):
                val = row[idx]
                if val is not None and str(val).strip():
                    return True
        return False
    data_rows = [r for r in data_rows if _row_has_data(r)]
    if not data_rows:
        return _err("No se encontraron filas de datos.")
    if len(data_rows) > MAX_ROWS:
        return _err(f"Excede el m\u00e1ximo de {MAX_ROWS} filas ({len(data_rows)} encontradas).")

    # Pre-scan: detectar emails duplicados dentro del archivo
    email_rows: dict[str, list[int]] = {}
    for scan_i, scan_rd in enumerate(data_rows, start=data_start + 1):
        e_i = col_idx.get("email")
        if e_i is not None and e_i < len(scan_rd) and scan_rd[e_i] is not None:
            e_val = str(scan_rd[e_i]).strip().lower()
            if e_val:
                email_rows.setdefault(e_val, []).append(scan_i)
    dup_emails = {e: filas for e, filas in email_rows.items() if len(filas) > 1}

    results: list[dict] = []
    created = 0
    errors = 0
    year = date.today().year

    for ri, rd in enumerate(data_rows, start=data_start + 1):
        def cv(col):
            idx = col_idx.get(col)
            if idx is None or idx >= len(rd) or rd[idx] is None:
                return ""
            return str(rd[idx]).strip()

        raw_name = cv("nombre_completo")
        raw_email = cv("email")
        raw_role = cv("rol")
        raw_team = cv("equipo")
        raw_date_val = rd[col_idx["fecha_ingreso"]] if col_idx.get("fecha_ingreso") is not None and col_idx["fecha_ingreso"] < len(rd) else None
        raw_pos = cv("puesto")
        raw_phone = cv("telefono")
        raw_emerg = cv("contacto_emergencia")
        raw_mgr_email = cv("gerente_email")

        entry = {"row": ri, "name": raw_name, "email": raw_email, "role": raw_role,
                 "team": raw_team, "manager": raw_mgr_email, "password": "", "vacation_days": 0,
                 "status": "VÁLIDO" if preview_only else "CREADO", "detail": ""}

        errs: list[str] = []
        if not raw_name:
            errs.append("Nombre completo es obligatorio")
        if not raw_email:
            errs.append("Email es obligatorio")
        elif "@" not in raw_email or "." not in raw_email:
            errs.append("Formato de email inválido")

        norm_role = _normalize_role(raw_role) if raw_role else None
        if not raw_role:
            errs.append("Rol es obligatorio")
        elif not norm_role:
            errs.append(f"Rol '{raw_role}' no válido. Use: Empleado o Gerente")

        if actor_role == "HR" and norm_role and norm_role not in VALID_ROLES:
            errs.append("RH solo puede importar Empleado o Gerente")

        if not raw_team:
            errs.append("Equipo es obligatorio")

        hire_date = _parse_date(raw_date_val) if raw_date_val else None
        if not raw_date_val:
            errs.append("Fecha de ingreso es obligatoria")
        elif not hire_date:
            errs.append(f"Fecha '{raw_date_val}' no válida (use AAAA-MM-DD)")

        if raw_email and raw_email.lower() in dup_emails:
            otras = [f for f in dup_emails[raw_email.lower()] if f != ri]
            if otras:
                errs.append(f"Email duplicado en el archivo (también en fila{'s' if len(otras) > 1 else ''} {', '.join(str(f) for f in otras)})")
            elif raw_email.lower() in seen:
                errs.append("Email duplicado en el archivo")

        if raw_email and not errs:
            if user_repo.get_by_email(raw_email):
                errs.append("Ya existe un usuario con este email")

        team: Team | None = None
        if raw_team and not any("Equipo" in e for e in errs):
            team, suggestion = _resolve_team(raw_team, all_teams)
            if not team:
                errs.append(f"Equipo '{raw_team}' no encontrado. Disponibles: {suggestion}")

        # Resolver gerente por email (opcional) o auto-asignar por equipo
        manager_user: User | None = None
        if raw_mgr_email:
            manager_user = user_repo.get_by_email(raw_mgr_email)
            if not manager_user:
                errs.append(f"Gerente '{raw_mgr_email}' no encontrado en el sistema")
            elif manager_user.role.value not in ("MANAGER", "ADMIN"):
                errs.append(f"'{raw_mgr_email}' no tiene rol de Gerente")
        elif team and norm_role == "EMPLOYEE":
            # Auto-asignar el manager del equipo si no se especificó gerente_email
            team_managers = [
                u for u in user_repo.list_all(role="MANAGER", team_id=str(team.id))
                if u.is_active
            ]
            if team_managers:
                manager_user = team_managers[0]

        if errs:
            entry["status"] = "ERROR"
            entry["detail"] = "; ".join(errs)
            errors += 1
            results.append(entry)
            continue

        # En modo preview solo calculamos vacaciones, no creamos nada
        if preview_only:
            vac_days = compute_vacation_days(hire_date, date.today()) if hire_date else 0
            seen.add(raw_email.lower())
            created += 1
            entry["vacation_days"] = vac_days
            mgr_note = f" (Gerente auto: {manager_user.full_name})" if manager_user and not raw_mgr_email else ""
            entry["detail"] = f"Listo para importar — {vac_days} días de vacaciones{mgr_note}"
            results.append(entry)
            continue

        sp = db.begin_nested()
        try:
            pwd = _gen_pwd()
            user = User(
                email=raw_email.lower(), full_name=raw_name,
                role=UserRole(norm_role), password_hash=hash_password(pwd),
                team_id=team.id if team else None, is_active=True,
                must_change_password=True, hire_date=hire_date,
                position=raw_pos or None, phone=raw_phone or None,
                emergency_contact=raw_emerg or None,
                manager_id=manager_user.id if manager_user else None,
            )
            user = user_repo.add(user)

            # Crear relación en tabla user_managers
            if manager_user:
                user_repo.set_managers(str(user.id), [str(manager_user.id)])

            vac_days = compute_vacation_days(hire_date, date.today())
            balance = VacationBalance(
                user_id=user.id, year=year,
                available_days=Decimal(str(vac_days)),
                used_days=Decimal("0"), carried_over_days=Decimal("0"),
            )
            balance_repo.add(balance)

            audit_repo.log(
                actor_user_id=actor_user_id, action="USER_CREATED_BULK",
                entity_type="user", entity_id=str(user.id),
                metadata={"email": raw_email, "role": norm_role, "full_name": raw_name,
                          "hire_date": str(hire_date), "vacation_days": vac_days,
                          "batch_id": batch_id},
            )

            sp.commit()
            seen.add(raw_email.lower())
            created += 1
            entry["password"] = pwd
            entry["vacation_days"] = vac_days
            mgr_note = f" | Gerente: {manager_user.full_name}" if manager_user else ""
            entry["detail"] = f"Creado correctamente — {vac_days} días de vacaciones{mgr_note}"

        except Exception as exc:
            sp.rollback()
            logger.exception("Error fila %d: %s", ri, exc)
            entry["status"] = "ERROR"
            entry["detail"] = f"Error interno: {exc}"
            errors += 1

        results.append(entry)

    return results, created, errors, batch_id


# ═══════════════════════════════════════════════════════════════════════════
# EXCEL DE RESULTADOS
# ═══════════════════════════════════════════════════════════════════════════

def generate_result_excel(results: list[dict]) -> bytes:
    wb = Workbook()
    bd = _brd()
    hdr_fill = PatternFill(start_color=_AZUL_OSCURO, end_color=_AZUL_OSCURO, fill_type="solid")
    hdr_font = Font(name="Calibri", size=11, bold=True, color=_BLANCO)
    ok_fill = PatternFill(start_color=_VERDE_CLARO, end_color=_VERDE_CLARO, fill_type="solid")
    err_fill = PatternFill(start_color=_ROJO_CLARO, end_color=_ROJO_CLARO, fill_type="solid")
    warn_fill = PatternFill(start_color=_AMARILLO, end_color=_AMARILLO, fill_type="solid")
    conf_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")

    ok_n = sum(1 for r in results if r.get("status") == "CREADO")
    err_n = sum(1 for r in results if r.get("status") == "ERROR")
    created_results = [r for r in results if r.get("status") == "CREADO" and r.get("password")]

    # ─────────────────────────────────────────────────────────────────────
    # HOJA 1: RESULTADOS (sin contraseñas)
    # ─────────────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resultados"
    ws.sheet_properties.tabColor = _VERDE

    # Banner
    ncols = 9
    for c in range(1, ncols + 1):
        ws.cell(row=1, column=c).fill = hdr_fill
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=1, column=1, value="SEEKOP — Resultados de Importación Masiva").font = Font(name="Calibri", size=16, bold=True, color=_BLANCO)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = ["Fila", "Nombre", "Email", "Rol", "Equipo", "Gerente", "Días Vacaciones", "Estado", "Detalle"]
    widths = [8, 30, 35, 15, 25, 30, 18, 12, 50]

    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=ci, value=h)
        c.font = hdr_font
        c.fill = PatternFill(start_color=_AZUL_MEDIO, end_color=_AZUL_MEDIO, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = bd
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = w

    for ri, res in enumerate(results, start=3):
        vals = [res.get("row",""), res.get("name",""), res.get("email",""), res.get("role",""),
                res.get("team",""), res.get("manager",""), res.get("vacation_days",""),
                res.get("status",""), res.get("detail","")]
        st = res.get("status", "")
        fill = ok_fill if st == "CREADO" else (err_fill if st == "ERROR" else warn_fill)
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.fill = fill
            c.border = bd
            c.font = Font(name="Calibri", size=10)

    # Resumen
    sr = len(results) + 5
    for c in range(1, 4):
        ws.cell(row=sr, column=c).fill = hdr_fill
    ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=3)
    ws.cell(row=sr, column=1, value="RESUMEN DE IMPORTACIÓN").font = Font(name="Calibri", size=12, bold=True, color=_BLANCO)
    ws.cell(row=sr, column=1).alignment = Alignment(horizontal="center")

    ws.cell(row=sr+1, column=1, value="Creados exitosamente:").font = Font(size=10, bold=True)
    ws.cell(row=sr+1, column=2, value=ok_n).font = Font(size=12, bold=True, color=_VERDE)
    ws.cell(row=sr+2, column=1, value="Con errores:").font = Font(size=10, bold=True)
    ws.cell(row=sr+2, column=2, value=err_n).font = Font(size=12, bold=True, color=_ROJO)
    ws.cell(row=sr+3, column=1, value="Total procesados:").font = Font(size=10, bold=True)
    ws.cell(row=sr+3, column=2, value=len(results)).font = Font(size=12, bold=True)

    if created_results:
        nr = sr + 5
        ws.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=ncols)
        ws.cell(row=nr, column=1, value="Las contraseñas temporales están en la hoja \"Credenciales\" de este archivo.").font = Font(bold=True, size=10, color=_AZUL_OSCURO)
        for cc in range(1, ncols + 1):
            ws.cell(row=nr, column=cc).fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

    ws.freeze_panes = "A3"

    # ─────────────────────────────────────────────────────────────────────
    # HOJA 2: CREDENCIALES (solo si hubo creaciones)
    # ─────────────────────────────────────────────────────────────────────
    if created_results:
        wc = wb.create_sheet("Credenciales")
        wc.sheet_properties.tabColor = "C00000"

        # Banner confidencial
        for c in range(1, 6):
            wc.cell(row=1, column=c).fill = conf_fill
        wc.merge_cells("A1:E1")
        wc.cell(row=1, column=1, value="CONFIDENCIAL — Credenciales Temporales").font = Font(name="Calibri", size=16, bold=True, color=_BLANCO)
        wc.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
        wc.row_dimensions[1].height = 35

        # Advertencia
        for c in range(1, 6):
            wc.cell(row=2, column=c).fill = PatternFill(start_color=_AMARILLO, end_color=_AMARILLO, fill_type="solid")
        wc.merge_cells("A2:E2")
        wc.cell(row=2, column=1, value="IMPORTANTE: Comparte cada contraseña SOLO con su empleado correspondiente. Elimina este archivo después de distribuirlas.").font = Font(bold=True, size=10, color=_ROJO)
        wc.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wc.row_dimensions[2].height = 35

        # Instrucciones de distribución
        wc.merge_cells("A3:E3")
        wc.cell(row=3, column=1, value="Instrucciones de distribución segura:").font = Font(size=10, bold=True, color=_AZUL_OSCURO)
        wc.row_dimensions[3].height = 20

        instrucciones = [
            "1. Envía a cada empleado SOLO su contraseña (no compartas la lista completa).",
            "2. Usa un canal seguro: en persona, mensaje directo o correo individual.",
            "3. Indica al empleado que debe cambiar su contraseña en el primer inicio de sesión.",
            "4. Elimina este archivo de tu computadora una vez distribuidas todas las contraseñas.",
            "5. NO imprimas esta hoja ni la dejes visible en tu escritorio.",
        ]
        for ii, instr in enumerate(instrucciones):
            row_i = 4 + ii
            wc.merge_cells(start_row=row_i, start_column=1, end_row=row_i, end_column=5)
            wc.cell(row=row_i, column=1, value=instr).font = Font(size=9, color="444444")
            for c in range(1, 6):
                wc.cell(row=row_i, column=c).fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        # Separador
        sep_row = 4 + len(instrucciones) + 1

        # Encabezados de la tabla
        cred_headers = ["#", "Nombre Completo", "Email", "Contraseña Temporal", "Primer Ingreso"]
        cred_widths = [6, 35, 35, 28, 20]
        for ci, h in enumerate(cred_headers, start=1):
            c = wc.cell(row=sep_row, column=ci, value=h)
            c.font = hdr_font
            c.fill = conf_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = bd
        for ci, w in enumerate(cred_widths, start=1):
            wc.column_dimensions[get_column_letter(ci)].width = w

        # Datos
        pwd_font = Font(name="Consolas", size=11, bold=True, color="C00000")
        alt_fill_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        alt_fill_2 = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        login_url = "seekop.com"

        for ci, cr in enumerate(created_results, start=1):
            row_n = sep_row + ci
            row_fill = alt_fill_1 if ci % 2 == 1 else alt_fill_2
            vals = [ci, cr.get("name",""), cr.get("email",""), cr.get("password",""), login_url]
            for vi, v in enumerate(vals, start=1):
                c = wc.cell(row=row_n, column=vi, value=v)
                c.border = bd
                c.fill = row_fill
                if vi == 4:
                    c.font = pwd_font
                elif vi == 1:
                    c.font = Font(size=10, bold=True)
                    c.alignment = Alignment(horizontal="center")
                else:
                    c.font = Font(name="Calibri", size=10)

        # Pie de nota
        foot_row = sep_row + len(created_results) + 2
        wc.merge_cells(start_row=foot_row, start_column=1, end_row=foot_row, end_column=5)
        wc.cell(row=foot_row, column=1, value=f"Total de credenciales: {len(created_results)} — Generado el {date.today().strftime('%d/%m/%Y')}").font = Font(size=9, italic=True, color="666666")

        foot2 = foot_row + 1
        for c in range(1, 6):
            wc.cell(row=foot2, column=c).fill = conf_fill
        wc.merge_cells(start_row=foot2, start_column=1, end_row=foot2, end_column=5)
        wc.cell(row=foot2, column=1, value="Elimina este archivo después de distribuir las contraseñas.").font = Font(size=10, bold=True, color=_BLANCO)
        wc.cell(row=foot2, column=1).alignment = Alignment(horizontal="center")

        wc.freeze_panes = f"A{sep_row + 1}"

        # Proteger hoja de credenciales con contraseña
        wc.protection.sheet = True
        wc.protection.password = "SeekopRH2026"
        wc.protection.enable()

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# ROLLBACK DE IMPORTACIÓN
# ═══════════════════════════════════════════════════════════════════════════

def rollback_import(
    batch_id: str, db: Session, actor_user_id: str,
) -> dict:
    """Desactiva todos los usuarios creados en un batch específico."""
    from sqlalchemy import select as sa_select
    from app.models.audit_log import AuditLog

    audit_repo = AuditRepository(db)
    user_repo = UserRepository(db)

    stmt = sa_select(AuditLog).where(
        AuditLog.action == "USER_CREATED_BULK",
        AuditLog.metadata_["batch_id"].astext == batch_id,
    )
    logs = list(db.execute(stmt).scalars().all())

    if not logs:
        return {"rolled_back": 0, "detail": "No se encontró el lote de importación o ya fue revertido."}

    deactivated = 0
    user_emails: list[str] = []

    for log_entry in logs:
        user_id = log_entry.entity_id
        user = user_repo.get_by_id(user_id)
        if user and user.is_active:
            user.is_active = False
            user_emails.append(user.email)
            deactivated += 1

    if deactivated > 0:
        audit_repo.log(
            actor_user_id=actor_user_id,
            action="BULK_IMPORT_ROLLBACK",
            entity_type="user",
            entity_id=batch_id,
            metadata={"batch_id": batch_id, "deactivated": deactivated, "emails": user_emails},
        )
        db.commit()

    return {
        "rolled_back": deactivated,
        "emails": user_emails,
        "detail": f"Se desactivaron {deactivated} usuario(s) del lote {batch_id[:8]}…",
    }
